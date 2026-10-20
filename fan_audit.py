#!/usr/bin/env python3
"""fan_audit.py — 用公开互动数据核对 X 账号的粉丝构成（只读、离线）。

输入（CSV，UTF-8，首行为表头）:
  following.csv   我关注的账号      handle,name,bio
  followers.csv   关注我的账号      handle,name,bio   （可以只是「已认证粉丝」）
  engagement.csv  公开互动统计      handle,likes,replies[,quotes,reposts]
  requests.csv    （可选）来求过互关的账号 handle[,note]
  keywords.txt    （可选）互关类关键词，每行一个

输出（--out 目录）:
  report.xlsx（装了 openpyxl 时）+ 每个分组一个 CSV + summary.md

脚本不联网、不调用任何 X 接口，也不会关注/取关任何人。
"""
import argparse, csv, os, sys, hashlib
from collections import Counter

DEFAULT_KEYWORDS = ["互关", "互粉", "互回", "回关", "有关必回", "必回", "浇水", "浇朋友",
                    "蓝V互", "交朋友", "交友", "诚信", "follow back", "followback", "mutual", "f4f"]


def norm(h):
    return (h or "").strip().lstrip("@").strip().lower()


def read_csv(path, required=("handle",)):
    if not path:
        return [], []
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    if rows:
        missing = [c for c in required if c not in rows[0]]
        if missing:
            sys.exit(f"{path} 缺少列: {missing}")
    return rows, (list(rows[0].keys()) if rows else [])


def load_people(path, label, issues):
    rows, _ = read_csv(path)
    out, blank, dup = {}, 0, 0
    for r in rows:
        h = norm(r.get("handle"))
        if not h:
            blank += 1
            continue
        if h in out:
            dup += 1
            continue
        out[h] = {"handle": "@" + (r.get("handle") or "").strip().lstrip("@").strip(),
                  "name": (r.get("name") or "").strip(),
                  "bio": " ".join((r.get("bio") or "").split())}
    issues.append(f"{label}: 读入 {len(rows)} 行，有效 {len(out)} 个，空用户名 {blank}，重复 {dup}")
    return out


def load_keywords(path):
    if not path:
        return DEFAULT_KEYWORDS
    kws = []
    with open(path, encoding="utf-8") as f:
        for line in f:
            s = line.strip()
            if s and not s.startswith("#"):
                kws.append(s.lower())
    return kws


def to_int(x):
    try:
        return int(float(x))
    except (TypeError, ValueError):
        return 0


def main():
    ap = argparse.ArgumentParser(description="X 粉丝构成核对（离线）")
    ap.add_argument("--following", required=True)
    ap.add_argument("--followers", required=True)
    ap.add_argument("--engagement", required=True)
    ap.add_argument("--requests", help="可选：来求过互关的账号列表")
    ap.add_argument("--keywords", help="可选：关键词文件，默认用内置词表")
    ap.add_argument("--out", default="output")
    ap.add_argument("--top", type=int, default=15, help="互动排行取前几名")
    ap.add_argument("--posts", default="?", help="互动数据覆盖了几条帖子（只用于写进报告）")
    ap.add_argument("--followers-label", default="粉丝", help="粉丝表口径，例如「已认证粉丝」")
    ap.add_argument("--anonymize", action="store_true", help="summary.md 里用 user_001 这类代号代替用户名")
    a = ap.parse_args()

    issues = []
    F = load_people(a.following, "关注表", issues)
    V = load_people(a.followers, a.followers_label + "表", issues)
    erows, _ = read_csv(a.engagement)
    E = {}
    for r in erows:
        h = norm(r.get("handle"))
        if not h:
            continue
        e = E.setdefault(h, {"likes": 0, "replies": 0, "quotes": 0, "reposts": 0, "handle": "@" + r["handle"].strip().lstrip("@")})
        for k in ("likes", "replies", "quotes", "reposts"):
            e[k] += to_int(r.get(k))
    issues.append(f"互动表: {len(erows)} 行，{len(E)} 个账号")
    req_rows, _ = read_csv(a.requests) if a.requests else ([], [])
    REQ = {norm(r.get("handle")) for r in req_rows if norm(r.get("handle"))}
    kws = load_keywords(a.keywords)

    def flag(h):
        p = V.get(h) or F.get(h) or {}
        text = (p.get("name", "") + " " + p.get("bio", "")).lower()
        return "|".join(k for k in kws if k in text)

    allh = set(F) | set(V) | set(E) | REQ
    rec = {}
    for h in allh:
        p = V.get(h) or F.get(h) or {}
        e = E.get(h, {})
        tot = sum(e.get(k, 0) for k in ("likes", "replies", "quotes", "reposts"))
        kw = flag(h)
        rec[h] = {"handle": p.get("handle") or e.get("handle") or "@" + h, "name": p.get("name", ""),
                  "follows_me": h in V, "i_follow": h in F,
                  "likes": e.get("likes", 0), "replies": e.get("replies", 0),
                  "quotes": e.get("quotes", 0), "reposts": e.get("reposts", 0), "total": tot,
                  "keyword_hits": kw, "asked_mutual": h in REQ,
                  "mutual_flag": bool(kw) or h in REQ, "bio": p.get("bio", "")}

    def srt(hs):
        return sorted((rec[h] for h in hs), key=lambda r: (-r["total"], r["handle"].lower()))

    mutual = set(F) & set(V)
    nfb = set(F) - set(V)          # 我关注、对方不在粉丝表
    fnf = set(V) - set(F)          # 对方关注我、我没关注
    zero_f = {h for h in V if rec[h]["total"] == 0}
    zero_m = {h for h in mutual if rec[h]["total"] == 0}
    zero_f_flag = {h for h in zero_f if rec[h]["mutual_flag"]}
    engagers = {h for h in E if rec[h]["total"] > 0}
    top = [r for r in srt(engagers)][: a.top]
    top_h = [norm(r["handle"]) for r in top]
    nfb_flag = {h for h in nfb if rec[h]["keyword_hits"]}
    fnf_eng = {h for h in fnf if rec[h]["total"] > 0}
    outside = {h for h in engagers if h not in V and h not in F}

    groups = [("全部账号", allh), ("粉丝", set(V)), ("我关注的", set(F)), ("互关", mutual),
              ("我关注但未回关", nfb), ("关注我但我未关注", fnf), ("粉丝中零互动", zero_f),
              ("互关中零互动", zero_m), ("零互动且有互关特征", zero_f_flag), ("未回关且简介含关键词", nfb_flag),
              ("未回关但有互动的粉丝", fnf_eng), (f"互动前{a.top}", top_h), ("有互动但不在两表", outside)]
    if REQ:
        groups.append(("求过互关的账号", REQ))

    os.makedirs(a.out, exist_ok=True)
    cols = ["handle", "name", "follows_me", "i_follow", "likes", "replies", "quotes", "reposts", "total",
            "keyword_hits", "asked_mutual", "mutual_flag", "bio"]
    for i, (title, hs) in enumerate(groups, 1):
        rows = [rec[h] for h in hs] if title.startswith("互动前") else srt(hs)
        with open(os.path.join(a.out, f"{i:02d}_{title}.csv"), "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=cols)
            w.writeheader()
            w.writerows(rows)
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for title, hs in groups:
            ws = wb.create_sheet(title[:31])
            ws.append(cols)
            rows = [rec[h] for h in hs] if title.startswith("互动前") else srt(hs)
            for r in rows:
                ws.append([r[c] for c in cols])
        ws = wb.create_sheet("数据说明")
        for x in issues:
            ws.append([x])
        wb.save(os.path.join(a.out, "report.xlsx"))
        xlsx_note = "report.xlsx"
    except ImportError:
        xlsx_note = "（没装 openpyxl，只输出了 CSV）"

    # summary.md
    alias = {}
    def show(h):
        h = norm(h)
        if not a.anonymize:
            return rec[h]["handle"]
        if h not in alias:
            alias[h] = f"user_{len(alias) + 1:03d}"
        return alias[h]

    def pct(x, y):
        return f"{x / y:.0%}" if y else "n/a"

    nV, nM = len(V), len(mutual)
    L = ["# 粉丝构成核对报告", ""]
    L.append(f"- 互动数据覆盖帖子数：{a.posts}；粉丝表口径：{a.followers_label}")
    L.append(f"- 关键词 {len(kws)} 个；有互动的账号 {len(engagers)} 个")
    L += ["", "## 基本数字", "", "| 项目 | 数量 |", "|---|---|",
          f"| 我关注的 | {len(F)} |", f"| {a.followers_label} | {nV} |", f"| 互关 | {nM} |",
          f"| 我关注、不在{a.followers_label}里 | {len(nfb)} |", f"| {a.followers_label}里我没关注的 | {len(fnf)} |",
          f"| {a.followers_label}中零公开互动 | {len(zero_f)}（{pct(len(zero_f), nV)}） |",
          f"| 互关中零公开互动 | {len(zero_m)}（{pct(len(zero_m), nM)}） |",
          f"| 零互动且有互关特征 | {len(zero_f_flag)} |",
          f"| 我关注、未回关、简介含关键词 | {len(nfb_flag)} |",
          f"| 我没回关、但有互动的粉丝 | {len(fnf_eng)} |",
          f"| 有互动、但两张表都不在 | {len(outside)} |"]
    if REQ:
        rq = [h for h in REQ]
        L += ["", "## 求过互关的账号", "",
              f"共 {len(rq)} 个：互关 {sum(h in mutual for h in rq)}，只关注我 {sum(h in fnf for h in rq)}，"
              f"只有我关注 {sum(h in nfb for h in rq)}，都不在 {sum(h not in V and h not in F for h in rq)}；"
              f"其中零公开互动 {sum(rec[h]['total'] == 0 for h in rq)} 个。"]
    L += ["", f"## 互动前 {a.top}", "", "| # | 账号 | 点赞 | 回复 | 合计 | 关系 |", "|---|---|---|---|---|---|"]
    for i, r in enumerate(top, 1):
        rel = "互关" if r["follows_me"] and r["i_follow"] else ("只有我关注" if r["i_follow"] else ("只关注我" if r["follows_me"] else "都不在表里"))
        L.append(f"| {i} | {show(r['handle'])} | {r['likes']} | {r['replies']} | {r['total']} | {rel} |")
    L.append("")
    L.append(f"前 {len(top)} 里互关的有 {sum(r['follows_me'] and r['i_follow'] for r in top)} 个。")
    L += ["", "## 数据说明", ""] + [f"- {x}" for x in issues]
    L += ["- 「零互动」只代表在统计的这些帖子上没有公开点赞/回复/引用/转帖，不代表对方没看过。",
          "- 关键词只是粗筛，会误判（尤其是自动翻译的简介），结论请人工复核。",
          f"- 输出文件：{xlsx_note} 和同目录下的 CSV。"]
    with open(os.path.join(a.out, "summary.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")
    print("\n".join(L))


if __name__ == "__main__":
    main()
